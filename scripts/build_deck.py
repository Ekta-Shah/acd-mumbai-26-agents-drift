"""Builds deck/drift-talk.pptx from deck/slides.yaml and data/results.json. Native charts, click-to-reveal animations, links."""
import copy
import json
import math
import os
import re

import yaml
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
Y = yaml.safe_load(open(os.path.join(ROOT, "deck", "slides.yaml")))
R = json.load(open(os.path.join(ROOT, "data", "results.json")))
ILLUSTRATIVE = R["meta"]["source"] != "live-run"


def _values():
    """Story numbers used in notes and slide text. They come from the results file, so a live run can never leave a stale number behind."""
    sm, daily = R["summary"], R["daily"]
    a, b = sm["no_fixes"], sm["all_fixes"]
    alarm = next((d["day"] for d in daily if d["drift_score_pct"] > R["meta"]["alarm_threshold_pct"]), None)
    live = R["meta"]["source"] == "live-run"
    return dict(lat_none=a["p95_latency_s"], lat_all=b["p95_latency_s"], cost_none="%.3f" % a["cost_usd_per_ticket"], cost_all="%.3f" % b["cost_usd_per_ticket"],
                cost_cents="%.1f" % (100 * a["cost_usd_per_ticket"]), replay_usd="%.2f" % b["replay_cost_usd_per_day"], replay_cents="%d" % round(100 * b["replay_cost_usd_per_day"]),
                breaches_none=a["above_limit_approved_without_human"], breaches_all=b["above_limit_approved_without_human"],
                sent_last=daily[-1]["sent_to_human"], hv=daily[-1]["high_value_tickets"], alarm_day=alarm, fix_max=max(R["runs"]["1111"]),
                run_word="the live run" if live else "the illustrative run",
                source_note="These numbers come from a live run of the demo you can run yourself." if live else "These numbers are illustrative until a live run replaces them.")


def _subst(x, v):
    if isinstance(x, str):
        for k, val in v.items():
            x = x.replace("{{%s}}" % k, str(val))
        assert "{{" not in x, x
        return x
    if isinstance(x, list):
        return [_subst(i, v) for i in x]
    if isinstance(x, dict):
        return {k: _subst(i, v) for k, i in x.items()}
    return x


Y = _subst(Y, _values())

TEMPLATE = os.path.join(ROOT, "template", "ACD Mumbai 2026 template.pptx")  # copy of the event template; slides are built on its master
BG, TX, SEC = "FFFFFF", "1F2937", "4B5563"
GRN, AMB, RED = "0B7A4B", "B45309", "C62828"  # same meanings as before, darkened to hold 4.5:1 on white
PANEL, LINE = "F3F4F9", "C3C9DC"
NAVY = "2A3990"  # template brand navy, used for titles
rgb = lambda h: RGBColor.from_string(h)
SW, SH = 10.0, 5.625  # template canvas
SX, OX = 0.75, -0.05  # plan coordinates (13.33 x 7.5) are mapped onto the template canvas
KY, YA = 0.615, 1.25  # content band 1.25 to 4.2 in, above the template's corner graphic
KF = 0.72
MX = 0.6
CW = 13.333 - 2 * MX
SHORT = Y["meta"]["short_title"]
BUTTON_Y = 4.4  # bottom-left strip stays clear of the corner graphic


def nx(x):
    return x * SX + OX


def ny(y):
    return YA + (y - 1.7) * KY


def fs(size):
    return max(12, round(size * KF * 2) / 2)


# ------------------------------------------------------------------ helpers
def style_run(r, size, color=TX, bold=False, font="Calibri", italic=False, raw=False):
    r.font.size = Pt(size if raw else fs(size))
    r.font.bold = bold
    r.font.italic = italic
    if font:
        r.font.name = font
    r.font.color.rgb = rgb(color)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, wrap=True, margin=0.08):
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin * SX)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for p in paras:
        par = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        par.alignment = p.get("align", PP_ALIGN.LEFT)
        par.space_after = Pt(p.get("after", 0) * KF)
        for seg in p["runs"] if "runs" in p else [p]:
            r = par.add_run()
            r.text = seg["text"]
            style_run(r, seg.get("size", p.get("size", 26)), seg.get("color", p.get("color", TX)), seg.get("bold", p.get("bold", False)),
                      seg.get("font", p.get("font", "Calibri")), seg.get("italic", False))


def geom(x, y, w, h, raw):
    if raw:
        return Inches(x), Inches(y), Inches(w), Inches(h)
    return Inches(nx(x)), Inches(ny(y)), Inches(w * SX), Inches(h * KY)


def textbox(shapes, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, wrap=True, raw=False):
    s = shapes.add_textbox(*geom(x, y, w, h, raw))
    fill_tf(s.text_frame, paras, anchor, wrap)
    if name:
        s.name = name
    return s


def box(shapes, x, y, w, h, paras=None, fill=PANEL, line=LINE, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, anchor=MSO_ANCHOR.MIDDLE, name=None, margin=0.15, raw=False):
    s = shapes.add_shape(shape, *geom(x, y, w, h, raw))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw * 0.7)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if paras:
        fill_tf(s.text_frame, paras, anchor, margin=margin)
    if name:
        s.name = name
    return s


def code(text, size=18, color=TX, after=0):
    return dict(text=text, size=size, color=color, font="Consolas", after=after)


def body(text, size=26, color=TX, bold=False, after=0, align=PP_ALIGN.LEFT):
    return dict(text=text, size=size, color=color, bold=bold, after=after, align=align)


# ------------------------------------------------------------------ template, layouts
prs = Presentation(TEMPLATE)
assert (prs.slide_width, prs.slide_height) == (9144000, 5143500)
LAY = {l.name: l for l in prs.slide_layouts}
TITLE_ONLY = LAY["TITLE_AND_BODY"]  # the layout the template slides use; carries the corner graphic and navy band
TPL = list(prs.slides)  # welcome, bio, topic, questions, thank you
FOOT_SRC = next(sh for sh in TPL[2].shapes if sh.has_text_frame and "Community Day" in sh.text_frame.text)._element  # navy band text from the template

SLIDES = []  # (spec, slide)
specs = list(Y["slides"]) + list(Y["appendix"])
for sp in specs:
    sl = prs.slides.add_slide(TITLE_ONLY)
    for ph in list(sl.placeholders):  # keep only the title placeholder
        if "TITLE" not in str(ph.placeholder_format.type):
            ph._element.getparent().remove(ph._element)
    SLIDES.append((sp, sl))
BY_ID = {str(sp["id"]): sl for sp, sl in SLIDES}


def set_title(sl, text, size=24, y=0.4, h=0.8):
    t = sl.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.34), Inches(y), Inches(8.6), Inches(h)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.text = ""
    r = tf.paragraphs[0].add_run()
    r.text = text
    style_run(r, size, NAVY, False, None, raw=True)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT


def stamp_footer(sl, label=None):
    """Event text from the template in the navy band, plus slide number and talk title on the left in white."""
    el = copy.deepcopy(FOOT_SRC)
    sl.shapes._spTree.append(el)
    for t in el.iter(qn("a:t")):
        t.text = t.text.replace("2025", "2026")
    s = sl.shapes.add_textbox(Inches(0.3), Inches(5.3), Inches(6), Inches(0.32))
    s.name = "Footer"
    tf = s.text_frame
    tf.margin_left = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    if label:
        r = p.add_run(); r.text = label + "  |  "; style_run(r, 12, "FFFFFF", raw=True)
    else:
        fld = parse_xml('<a:fld %s id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum"><a:rPr lang="en-US" sz="1200"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:latin typeface="Calibri"/></a:rPr><a:t>1</a:t></a:fld>' % nsdecls("a"))
        p._p.append(fld)
    r = p.add_run(); r.text = ("  |  " if not label else "") + SHORT; style_run(r, 12, "FFFFFF", raw=True)


def tag(sl):
    if ILLUSTRATIVE:
        box(sl.shapes, 7.75, 0.07, 1.9, 0.3, [dict(text="ILLUSTRATIVE", size=18, bold=True, color=AMB, align=PP_ALIGN.CENTER)], fill=None, line=AMB, lw=2, name="Illustrative tag", raw=True, margin=0.02)


def link_button(sl, x, y, w, text, url=None, target=None, h=0.6, size=20):
    s = box(sl.shapes, nx(x) , BUTTON_Y, w * SX, 0.4, [dict(text=text, size=size, bold=True, color=NAVY, align=PP_ALIGN.CENTER)], fill=PANEL, line=NAVY, lw=2, name="Link: " + text, raw=True, margin=0.04)
    if url:
        s.click_action.hyperlink.address = url
    if target is not None:
        s.click_action.target_slide = target
    return s


def back_button(sl):
    link_button(sl, MX, 6.35, 2.6, "Back to hub", target=BY_ID["5"])


# ------------------------------------------------------------------ animation + transition XML
def add_reveals(sl, groups):
    """groups: list of lists of shapes. One click per group; shapes in a group appear together (fade)."""
    ctn = [2]

    def nid():
        ctn[0] += 1
        return ctn[0]

    clicks = ""
    blds = ""
    for g in groups:
        pars = ""
        for j, shp in enumerate(g):
            spid = shp.shape_id
            node = "clickEffect" if j == 0 else "withEffect"
            grp = ' grpId="0"' if shp._element.tag == qn("p:sp") else ""
            pars += ('<p:par><p:cTn id="%d" presetID="10" presetClass="entr" presetSubtype="0" fill="hold"' + grp + ' nodeType="%s"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                     '<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                     '<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="%d" dur="300"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:animEffect>'
                     '</p:childTnLst></p:cTn></p:par>') % (nid(), node, nid(), spid, nid(), spid)
            if shp._element.tag == qn("p:sp"):
                blds += '<p:bldP spid="%d" grpId="0" animBg="1"/>' % spid
        a, b = nid(), nid()
        clicks += ('<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   '<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
                   '</p:childTnLst></p:cTn></p:par>') % (a, b, pars)
    xml = ('<p:timing %s><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>%s</p:childTnLst></p:cTn>'
           '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>%s</p:bldLst></p:timing>') % (nsdecls("p"), clicks, blds)
    sl._element.append(parse_xml(xml))


def fade_transition(sl):
    el = parse_xml('<p:transition %s spd="med"><p:fade/></p:transition>' % nsdecls("p"))
    clr = sl._element.find(qn("p:clrMapOvr"))
    clr.addnext(el)


def group_of(sl, build):
    """Build shapes inside a group so one animation reveals them all. build(shapes) adds shapes."""
    g = sl.shapes.add_group_shape()
    build(g.shapes)
    return g


# ------------------------------------------------------------------ native charts
def cx(tag_, **attrs):
    return parse_xml('<c:%s %s %s/>' % (tag_, nsdecls("c"), " ".join('%s="%s"' % kv for kv in attrs.items())))


def line_chart(sl, x, y, w, h, cats, series, title, ymin=None, ymax=None, yfmt="General", major=None, skip=5, name="Chart"):
    """series: list of dict(name, values, color, dash, label, points={index: text}) ; direct labels, no legend."""
    cd = CategoryChartData()
    cd.categories = cats
    for s in series:
        cd.add_series(s["name"], s["values"])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(nx(x)), Inches(ny(y)), Inches(w * SX), Inches(h * KY), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    for p in ch.chart_title.text_frame.paragraphs:
        for r in p.runs:
            style_run(r, 22, TX, True)
    ch.chart_title.include_in_layout = False
    ch.font.size = Pt(13)
    ch.font.color.rgb = rgb(TX)
    ch.font.name = "Calibri"
    va, ca = ch.value_axis, ch.category_axis
    if ymin is not None:
        va.minimum_scale = ymin
    if ymax is not None:
        va.maximum_scale = ymax
    if major:
        va.major_unit = major
    va.tick_labels.font.size = Pt(13); va.tick_labels.font.color.rgb = rgb(SEC); va.tick_labels.number_format = yfmt; va.tick_labels.number_format_is_linked = False
    ca.tick_labels.font.size = Pt(13); ca.tick_labels.font.color.rgb = rgb(SEC)
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(LINE)
    va.format.line.fill.background()
    ca.format.line.color.rgb = rgb(SEC)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    cax = ca._element
    for t in ("c:tickLblSkip", "c:tickMarkSkip"):
        for old in cax.findall(qn(t)):
            cax.remove(old)
    nm = cax.find(qn("c:noMultiLvlLbl"))
    for el in (cx("tickLblSkip", val=skip), cx("tickMarkSkip", val=skip)):
        if nm is not None:
            nm.addprevious(el)
        else:
            cax.append(el)
    plot = ch.plots[0]
    for s, spec in zip(plot.series, series):
        s.smooth = False
        s.format.line.color.rgb = rgb(spec["color"])
        s.format.line.width = Pt(spec.get("width", 4.5) * 0.7)
        if spec.get("dash"):
            s.format.line.dash_style = MSO_LINE.DASH
        s.marker.style = XL_MARKER_STYLE.NONE
        labels = dict(spec.get("points", {}))
        if spec.get("label"):
            labels[len(spec["values"]) - 1] = spec["label"]
        for idx, text in labels.items():
            pt = s.points[idx]
            if idx in spec.get("points", {}):
                pt.marker.style = XL_MARKER_STYLE.CIRCLE
                pt.marker.size = 8
                pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = rgb(BG)
                pt.marker.format.line.color.rgb = rgb(spec["color"]); pt.marker.format.line.width = Pt(2)
            dl = pt.data_label
            dl.has_text_frame = True
            dl.text_frame.text = text
            for p in dl.text_frame.paragraphs:
                for r in p.runs:
                    style_run(r, 18, spec["color"], True)
            dl.position = spec.get("pos") or (XL_LABEL_POSITION.RIGHT if idx == len(spec["values"]) - 1 and idx not in spec.get("points", {}) else XL_LABEL_POSITION.ABOVE)
    cs = ch._chartSpace
    sp = parse_xml('<c:spPr %s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>' % nsdecls("c"))
    cs.find(qn("c:chart")).addnext(sp)
    pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    pa.append(parse_xml('<c:spPr %s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></c:spPr>' % nsdecls("c")))
    return gf


DAYS = [str(d["day"]) for d in R["daily"]]


def ytop(*series):
    """Axis top for drift charts: 50 percent, or higher when a real run drifts further."""
    return max(50, math.ceil((max(max(s) for s in series) + 5) / 10) * 10)


def ymajor(top):
    return 10 if top <= 60 else 20
col = lambda k: [d[k] for d in R["daily"]]


# ------------------------------------------------------------------ slide kinds
def k_title(sp, sl):
    set_title(sl, sp["title"], size=36, y=1.35, h=0.7)
    t = sl.shapes.add_textbox(Inches(0.34), Inches(2.15), Inches(8.6), Inches(0.8))
    fill_tf(t.text_frame, [body(sp["subtitle"], 30, SEC)])
    box(sl.shapes, 0.4, 3.1, 0.9, 0.05, fill=AMB, line=None, shape=MSO_SHAPE.RECTANGLE, raw=True)
    textbox(sl.shapes, 0.34, 3.3, 8.6, 0.4, [body(Y["meta"]["speaker"], 26)], raw=True)
    textbox(sl.shapes, 0.34, 3.7, 8.6, 0.4, [body(Y["meta"]["event"] + ", 17 October 2026", 26, SEC)], raw=True)


def card(shapes, x, y, w, h, c, drift=False):
    box(shapes, x, y, w, h, fill=PANEL, line=AMB if drift else LINE, lw=3 if drift else 1.5)
    textbox(shapes, x + 0.2, y + 0.15, w - 0.4, 0.5, [body(c["day"], 26, SEC, True)])
    textbox(shapes, x + 0.2, y + 0.75, w - 0.4, 0.6, [body(c["ticket"], 28, TX, True)])
    textbox(shapes, x + 0.2, y + 1.55, w - 0.4, 1.4, [code(c["trace"], 20)])
    textbox(shapes, x + 0.2, y + 3.2, w - 0.4, 0.8, [dict(runs=[dict(text=c["outcome"], size=32, bold=True, color=AMB if drift else TX)] + ([dict(text="   DRIFT", size=24, bold=True, color=AMB)] if drift else []))])


def k_two_cards(sp, sl):
    set_title(sl, sp["title"])
    if "e1" in R:  # measured outcome of ticket T-4471 in the run replaces the example wording
        sp = dict(sp, left=dict(sp["left"], trace=R["e1"]["day1"]["trace"], outcome=R["e1"]["day1"]["outcome"]),
                  right=dict(sp["right"], trace=R["e1"]["day30"]["trace"], outcome=R["e1"]["day30"]["outcome"]))
    card(sl.shapes, MX, 2.0, 5.9, 4.3, sp["left"])
    g = group_of(sl, lambda shp: card(shp, MX + 6.23, 2.0, 5.9, 4.3, sp["right"], drift=True))
    g.name = "Day 30 card"
    add_reveals(sl, [[g]])


def k_panels(sp, sl):
    set_title(sl, sp["title"])
    s = R["summary"]["no_fixes"]
    vals = [f"{s['p95_latency_s']} s", f"{R['daily'][0]['error_rate_pct']:.1f}%", f"${s['cost_usd_per_ticket']:.3f}"]
    for i, (p, v) in enumerate(zip(sp["panels"], vals)):
        x = MX + i * 4.12
        box(sl.shapes, x, 2.0, 3.9, 2.7, [body(p["label"], 26, SEC, True, 4), body(v, 48, GRN, True, 4), body("OK", 28, GRN, True)], fill=PANEL, line=GRN, lw=2)
    last = R["daily"][-1]
    h = sp["hidden"]

    def hid(shp):
        box(shp, MX, 5.0, CW, 1.5, fill=PANEL, line=AMB, lw=3)
        textbox(shp, MX + 0.3, 5.2, CW - 0.6, 1.1, [dict(runs=[dict(text=h["label"] + ": ", size=40, color=TX, bold=True), dict(text=f"{last['sent_to_human']} of {last['high_value_tickets']}", size=48, color=AMB, bold=True), dict(text="   DRIFT", size=24, color=AMB, bold=True)])], anchor=MSO_ANCHOR.MIDDLE)
    g = group_of(sl, hid)
    g.name = "Hidden panel"
    add_reveals(sl, [[g]])


def k_statement(sp, sl):
    set_title(sl, sp["title"])
    textbox(sl.shapes, MX, 2.5, CW, 1.4, [body(sp["lines"][0], 40, TX, True)])
    textbox(sl.shapes, MX, 4.4, 5.5, 0.9, [body(sp["lines"][1], 44, SEC, True)])
    textbox(sl.shapes, MX + 6.0, 4.4, 6.0, 0.9, [body(sp["lines"][2], 44, AMB, True)])


def k_hub(sp, sl):
    set_title(sl, sp["title"])
    for i, t in enumerate(sp["tiles"]):
        x, y = MX + (i % 2) * 6.2, 2.0 + (i // 2) * 2.35
        s = box(sl.shapes, x, y, 5.93, 2.1, [body(t["name"], 44, TX, True, 2), body(t["sub"], 28, SEC)], fill=PANEL, line=AMB, lw=3, name="Tile: " + t["name"])
        s.click_action.target_slide = BY_ID[str(t["target"])]


def k_trace_reveal(sp, sl):
    set_title(sl, sp["title"])
    box(sl.shapes, MX, 2.1, CW, 2.5, fill=PANEL, line=LINE)
    for i, ln in enumerate(sp["lines"]):
        textbox(sl.shapes, MX + 0.25, 2.3 + i * 1.05, CW - 0.5, 0.9, [code(ln, 20)], anchor=MSO_ANCHOR.MIDDLE)
    hl = box(sl.shapes, MX + 0.15, 2.3 + sp["highlight"] * 1.05, CW - 0.3, 0.95, fill=None, line=AMB, lw=4, shape=MSO_SHAPE.RECTANGLE, name="Highlight")
    co = textbox(sl.shapes, MX, 4.95, CW, 0.8, [body(sp["callout"], 36, AMB, True)], name="Callout")
    add_reveals(sl, [[hl, co]])
    if sp.get("back"):
        back_button(sl)


def k_json_compare(sp, sl):
    set_title(sl, sp["title"])
    for i, (lab, txt) in enumerate(((sp["before_label"], sp["before"]), (sp["after_label"], sp["after"]))):
        x = MX + i * 6.23
        box(sl.shapes, x, 2.0, 5.9, 2.2, fill=PANEL, line=LINE)
        textbox(sl.shapes, x + 0.2, 2.1, 5.5, 0.5, [body(lab, 24, SEC, True)])
        textbox(sl.shapes, x + 0.2, 2.7, 5.5, 1.4, [code(txt, 20)])
    ag = box(sl.shapes, MX, 4.55, CW, 0.9, [code(sp["agent"], 22)], fill=PANEL, line=RED, lw=3, name="Agent line")
    bd = textbox(sl.shapes, MX, 5.55, 5, 0.7, [dict(runs=[dict(text="! ", size=32, bold=True, color=RED), dict(text=sp["badge"], size=32, bold=True, color=RED)])], name="Breach badge")
    add_reveals(sl, [[ag, bd]])
    if sp.get("back"):
        back_button(sl)


def k_steps_trace(sp, sl):
    set_title(sl, sp["title"])
    groups = []
    for i, st in enumerate(sp["steps"]):
        y = 2.0 + i * 1.4

        def row(shp, st=st, y=y):
            box(shp, MX, y, CW, 1.2, fill=PANEL, line=RED if st.get("badge") else LINE, lw=3 if st.get("badge") else 1.5)
            textbox(shp, MX + 0.2, y + 0.15, 2.0, 0.9, [body(st["label"], 28, SEC, True)], anchor=MSO_ANCHOR.MIDDLE)
            textbox(shp, MX + 2.3, y + 0.15, 7.7, 0.9, [code(l, 20) for l in st["lines"]], anchor=MSO_ANCHOR.MIDDLE)
            if st.get("badge"):
                textbox(shp, MX + 10.2, y + 0.3, 1.8, 0.6, [body("! " + st["badge"], 26, RED, True)], anchor=MSO_ANCHOR.MIDDLE)
        g = group_of(sl, row)
        g.name = "Step " + st["label"]
        groups.append([g])
    add_reveals(sl, groups)
    if sp.get("back"):
        back_button(sl)


def k_handoff(sp, sl):
    set_title(sl, sp["title"])
    groups = []
    ys = [1.95, 3.85, 5.0]
    hs = [1.75, 0.95, 0.8]
    for h, y, hh in zip(sp["hops"], ys, hs):
        def row(shp, h=h, y=y, hh=hh):
            box(shp, MX, y, 8.4, hh, fill=PANEL, line=LINE)
            textbox(shp, MX + 0.15, y + 0.05, 8.1, hh - 0.1, [dict(text=h["label"], size=20, color=SEC, bold=True, font="Consolas", after=2), code(h["text"], 18)], anchor=MSO_ANCHOR.MIDDLE)
            if h["lost"]:
                textbox(shp, MX + 8.65, y, 3.5, hh, [body(h["lost"], 24, AMB, True)], anchor=MSO_ANCHOR.MIDDLE)
        g = group_of(sl, row)
        g.name = "Hop"
        groups.append([g])
    add_reveals(sl, groups)
    if sp.get("back"):
        back_button(sl)


def k_chart_retention(sp, sl):
    set_title(sl, sp["title"])
    ns = list(range(1, 7))
    s = [dict(name="95 percent kept per handoff", values=[round(100 * 0.95 ** n, 1) for n in ns], color=TX, label="95%: 74% left", pos=XL_LABEL_POSITION.ABOVE),
         dict(name="90 percent kept per handoff", values=[round(100 * 0.90 ** n, 1) for n in ns], color=SEC, dash=True, label="90%: 53% left", pos=XL_LABEL_POSITION.BELOW)]
    line_chart(sl, MX, 1.7, CW, 4.6, [str(n) for n in ns], s, sp["chart_title"] + " (handoffs 1 to 6)", ymin=0, ymax=100, yfmt='0"%"', major=25, skip=1, name="Retention chart")
    link_button(sl, MX, 6.4, 5.2, sp["link"]["text"], url=sp["link"]["url"])


def arrow(shapes, x, y, w=0.45, h=0.4, color=SEC):
    return box(shapes, x, y, w, h, fill=color, line=None, shape=MSO_SHAPE.RIGHT_ARROW)


def k_architecture(sp, sl):
    set_title(sl, sp["title"])

    def l1(shp):
        box(shp, MX, 1.95, 8.2, 2.55, fill=None, line=TX, lw=2.5, anchor=MSO_ANCHOR.TOP)
        textbox(shp, MX + 0.15, 2.0, 7.9, 0.5, [body("AgentCore Runtime  |  Strands Agents SDK", 22, SEC, True)])
        for i, n in enumerate(("Intake", "Policy", "Resolution")):
            box(shp, MX + 0.3 + i * 2.65, 2.75, 2.2, 1.4, [body(n, 28, TX, True, align=PP_ALIGN.CENTER)], fill=PANEL, line=TX, lw=2)
        for i in range(2):
            arrow(shp, MX + 2.55 + i * 2.65, 3.3, 0.4, 0.3)
    g1 = group_of(sl, l1)

    def l2(shp):
        box(shp, MX, 4.8, 3.95, 1.1, [body("Amazon Bedrock", 26, TX, True)], fill=PANEL, line=LINE)
        box(shp, MX + 4.25, 4.8, 3.95, 1.1, [body("Bedrock Guardrails", 26, TX, True)], fill=PANEL, line=LINE)
    g2 = group_of(sl, l2)
    g3 = group_of(sl, lambda shp: box(shp, 9.2, 1.95, 3.53, 1.3, [body("AgentCore Memory", 26, TX, True)], fill=PANEL, line=LINE))
    g4 = group_of(sl, lambda shp: box(shp, 9.2, 3.45, 3.53, 1.3, [body("Replay harness", 26, TX, True)], fill=PANEL, line=AMB, lw=2.5))
    g5 = group_of(sl, lambda shp: box(shp, 9.2, 4.95, 3.53, 1.3, [body("Amazon CloudWatch", 26, TX, True)], fill=PANEL, line=AMB, lw=2.5))
    add_reveals(sl, [[g1], [g2], [g3], [g4], [g5]])
    if sp.get("link"):
        link_button(sl, MX, 6.4, 4.2, sp["link"]["text"], url=sp["link"]["url"])


def k_metrics_vs_drift(sp, sl):
    set_title(sl, sp["title"])
    box(sl.shapes, MX, 2.0, 5.6, 3.4, fill=PANEL, line=GRN, lw=2, anchor=MSO_ANCHOR.TOP)
    textbox(sl.shapes, MX + 0.25, 2.1, 5.1, 0.5, [body("Dashboard", 26, SEC, True)])
    textbox(sl.shapes, MX + 0.25, 2.75, 5.1, 2.5, [body(t, 40, GRN, True, 6) for t in sp["left_items"]])
    box(sl.shapes, 6.5, 2.0, 6.23, 3.4, fill=PANEL, line=AMB, lw=3, anchor=MSO_ANCHOR.TOP)
    textbox(sl.shapes, 6.75, 2.1, 5.7, 0.5, [body("Drift score", 26, SEC, True)])
    textbox(sl.shapes, 6.75, 2.8, 5.7, 2.4, [code(sp["formula"], 22, AMB)])
    gs = textbox(sl.shapes, MX, 5.75, CW, 0.8, [body(sp["golden"], 32, TX, True)], name="Golden set")
    add_reveals(sl, [[gs]])


def k_demo_green(sp, sl):
    set_title(sl, sp["title"])
    spec = [("p95_latency_s", "Latency stayed flat", "p95 latency", 0, 3, 1, "0.0"), ("error_rate_pct", "Errors stayed at zero", "errors", 0, 1, 0.5, "0.0"), ("cost_usd_per_ticket", "Cost stayed flat", "cost", 0, 0.02, 0.01, "0.000")]
    for i, (k, title, lab, lo, hi, mj, fmt) in enumerate(spec):
        line_chart(sl, MX + i * 4.12, 1.75, 3.95, 3.7, DAYS, [dict(name=lab, values=col(k), color=GRN, label=None)], title, ymin=lo, ymax=hi, yfmt=fmt, major=mj, skip=9, name="Chart " + lab)
    textbox(sl.shapes, MX, 5.65, CW, 0.7, [body("Drag the day scrubber from 1 to 30", 28, SEC)])
    link_button(sl, MX, 6.3, 5.6, sp["link"]["text"], url=sp["link"]["url"])


def k_demo_catch(sp, sl):
    set_title(sl, sp["title"])
    pts = {8: "Day 9", 13: "Day 14", 23: "Day 24"}
    line_chart(sl, MX, 1.7, CW, 4.1, DAYS, [dict(name="Drift score", values=col("drift_score_pct"), color=AMB, points=pts, label="drift score")], "The drift score rose on days 9, 14 and 24", ymin=0, ymax=ytop(col("drift_score_pct")), yfmt='0"%"', major=ymajor(ytop(col("drift_score_pct"))), skip=1, name="Drift chart")
    names = ["Day 9: memory", "Day 14: goal, handoff", "Day 24: tool"]
    for i, (lk, nm) in enumerate(zip(sp["links"], names)):
        link_button(sl, MX + i * 3.0, 6.25, 2.85, nm, url=lk["url"], size=18)


def k_demo_fix(sp, sl):
    set_title(sl, sp["title"])
    s = [dict(name="no fixes", values=R["runs"]["0000"], color=AMB, dash=True, label="no fixes: %s%%" % R["runs"]["0000"][-1]), dict(name="all four patterns", values=R["runs"]["1111"], color=AMB, width=6, label="all four")]
    line_chart(sl, MX, 1.7, CW, 4.4, DAYS, s, f"All four patterns kept drift at or below {max(R['runs']['1111'])} percent", ymin=0, ymax=ytop(R["runs"]["0000"]), yfmt='0"%"', major=ymajor(ytop(R["runs"]["0000"])), skip=1, name="Fix chart")
    link_button(sl, MX, 6.3, 5.6, sp["link"]["text"], url=sp["link"]["url"])


def k_pattern_pin(sp, sl):
    set_title(sl, sp["title"])
    box(sl.shapes, MX, 2.0, 5.9, 1.8, fill=PANEL, line=AMB, lw=3, anchor=MSO_ANCHOR.TOP)
    textbox(sl.shapes, MX + 0.2, 2.05, 5.5, 0.5, [body("BEFORE", 24, AMB, True)])
    textbox(sl.shapes, MX + 0.2, 2.65, 5.5, 1.1, [code(sp["before"], 20)])

    def aft(shp):
        box(shp, MX + 6.23, 2.0, 5.9, 1.8, fill=PANEL, line=GRN, lw=3, anchor=MSO_ANCHOR.TOP)
        textbox(shp, MX + 6.43, 2.05, 5.5, 0.5, [body("AFTER", 24, GRN, True)])
        textbox(shp, MX + 6.43, 2.65, 5.5, 1.1, [code(sp["after"], 20)])
    g = group_of(sl, aft)
    textbox(sl.shapes, MX, 4.35, CW, 2.2, [code(m, 20, SEC, 10) for m in sp["mapping"]])
    add_reveals(sl, [[g]])


def k_pattern_memory(sp, sl):
    set_title(sl, sp["title"])
    for i, r in enumerate(sp["rules"]):
        box(sl.shapes, MX, 2.0 + i * 0.95, 7.0, 0.8, [body(f"{i + 1}  {r}", 30, TX, True)], fill=PANEL, line=LINE)
    box(sl.shapes, MX, 5.0, CW, 0.9, [dict(runs=[dict(text=sp["trace"], size=20, font="Consolas", color=TX), dict(text="   REJECTED", size=24, bold=True, color=GRN)])], fill=PANEL, line=GRN, lw=3)
    textbox(sl.shapes, MX, 6.1, CW, 0.6, [code(sp["mapping"], 20, SEC)])


def k_pattern_rules(sp, sl):
    set_title(sl, sp["title"])
    box(sl.shapes, MX, 2.0, 6.7, 2.4, [code(l, 24) for l in sp["code"]], fill=PANEL, line=GRN, lw=3)
    box(sl.shapes, MX + 7.0, 2.0, 5.13, 2.4, [body("Bedrock Guardrails", 24, SEC, True, 6), body(sp["guardrail"], 28)], fill=PANEL, line=LINE, anchor=MSO_ANCHOR.MIDDLE)
    textbox(sl.shapes, MX, 4.9, CW, 1.0, [body(sp["footer"], 36, TX, True)])


def k_pattern_replay(sp, sl):
    set_title(sl, sp["title"])
    n = len(sp["flow"])
    w = (CW - (n - 1) * 0.55) / n
    for i, t in enumerate(sp["flow"]):
        box(sl.shapes, MX + i * (w + 0.55), 1.75, w, 0.8, [body(t, 22, TX, True, align=PP_ALIGN.CENTER)], fill=PANEL, line=AMB if i >= 2 else LINE, lw=2)
        if i < n - 1:
            arrow(sl.shapes, MX + i * (w + 0.55) + w + 0.07, 1.95, 0.4, 0.4)
    thr = [R["meta"]["alarm_threshold_pct"]] * len(DAYS)
    first = next(d["day"] for d in R["daily"] if d["drift_score_pct"] > R["meta"]["alarm_threshold_pct"])
    line_chart(sl, MX, 2.7, CW, 3.5, DAYS, [dict(name="Drift score", values=col("drift_score_pct"), color=AMB, label="drift score"), dict(name="Alarm", values=thr, color=SEC, dash=True, width=3, label="alarm %s%%" % R["meta"]["alarm_threshold_pct"])],
               f"The score crossed the {R['meta']['alarm_threshold_pct']} percent alarm on day {first}", ymin=0, ymax=ytop(col("drift_score_pct")), yfmt='0"%"', major=ymajor(ytop(col("drift_score_pct"))), skip=1, name="Alarm chart")
    al = box(sl.shapes, 1.6 + 8.6 * (first - 1) / 29, 3.5, 3.8, 0.8, [body(f"ALARM: day {first}", 32, AMB, True, align=PP_ALIGN.CENTER)], fill=BG, line=AMB, lw=3, name="Alarm label")
    textbox(sl.shapes, MX, 6.35, CW, 0.6, [code(sp["mapping"], 20, SEC)])
    add_reveals(sl, [[al]])


def k_results_table(sp, sl):
    set_title(sl, sp["title"])
    a, b = R["summary"]["no_fixes"], R["summary"]["all_fixes"]
    rows = [("Measure, days 1 to 30", "No fixes", "All four"),
            ("Approved over limit, no human", str(a["above_limit_approved_without_human"]), str(b["above_limit_approved_without_human"])),
            ("p95 latency", f"{a['p95_latency_s']} s", f"{b['p95_latency_s']} s"),
            ("Cost per ticket", f"${a['cost_usd_per_ticket']:.3f}", f"${b['cost_usd_per_ticket']:.3f}"),
            ("Replay cost per day", "none" if not a["replay_cost_usd_per_day"] else f"${a['replay_cost_usd_per_day']:.2f}", f"about ${b['replay_cost_usd_per_day']:.2f}")]
    gf = sl.shapes.add_table(len(rows), 3, Inches(nx(MX)), Inches(ny(1.8)), Inches(CW * SX), Inches(0.75 * len(rows) * KY))
    gf.name = "Results table"
    t = gf.table
    for ci, wd in enumerate((6.4, 2.7, 3.03)):
        t.columns[ci].width = Inches(wd * SX)
    for ri, r in enumerate(rows):
        t.rows[ri].height = Inches(0.75 * KY)
        for ci, v in enumerate(r):
            c = t.cell(ri, ci)
            c.fill.solid(); c.fill.fore_color.rgb = rgb(LINE if ri == 0 else PANEL)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = ""
            run = tf.paragraphs[0].add_run()
            color, text = TX, v
            if ri == 1 and ci > 0:
                bad = int(v) > 0
                color, text = (RED, f"{v}  BREACH") if bad else (GRN, f"{v}  OK")
            run.text = text
            style_run(run, 26, color if ri else SEC, bold=(ri == 0 or (ri == 1 and ci > 0)))
    diff_c = b["cost_usd_per_ticket"] - a["cost_usd_per_ticket"]
    diff_l = b["p95_latency_s"] - a["p95_latency_s"]
    mw = lambda v, unit: f"{abs(v):.{3 if unit == 'USD' else 1}f} {unit} " + ("more" if v >= 0 else "less")
    textbox(sl.shapes, MX, 5.75, CW, 0.7, [body(f"Fixes cost {mw(diff_c, 'USD')} per ticket and {mw(diff_l, 's')} latency.", 28, TX, True)])
    link_button(sl, MX, 6.35, 5.2, sp["link"]["text"], url=sp["link"]["url"])


def k_checklist(sp, sl):
    set_title(sl, sp["title"])
    for i, it in enumerate(sp["items"]):
        y = 1.95 + i * 0.95
        box(sl.shapes, MX, y, 0.7, 0.7, [body(str(i + 1), 30, BG, True, align=PP_ALIGN.CENTER)], fill=TX, line=None, shape=MSO_SHAPE.OVAL, margin=0.02)
        textbox(sl.shapes, MX + 0.95, y, 7.6, 0.75, [body(it, 32, TX, True)], anchor=MSO_ANCHOR.MIDDLE)
    box(sl.shapes, 9.55, 2.0, 3.18, 3.18, [body("QR", 54, SEC, True, align=PP_ALIGN.CENTER)], fill=None, line=SEC, lw=3, shape=MSO_SHAPE.RECTANGLE, name="QR placeholder")
    textbox(sl.shapes, 9.55, 5.3, 3.18, 0.9, [body(sp["qr_label"], 22, SEC, align=PP_ALIGN.CENTER)])


def k_closing(sp, sl):
    set_title(sl, sp["title"], size=30, y=1.45, h=1.0)
    t = sl.shapes.add_textbox(Inches(0.34), Inches(2.6), Inches(8.6), Inches(0.8))
    fill_tf(t.text_frame, [body(sp["sub"], 40, SEC)])


def k_appendix_text(sp, sl):
    set_title(sl, sp["title"])
    for i, it in enumerate(sp["items"]):
        box(sl.shapes, MX, 1.95 + i * 0.9, 5.4, 0.75, [body(it, 28, TX, True)], fill=PANEL, line=LINE)
    ba = sp["before_after"]
    box(sl.shapes, 6.3, 1.95, 6.43, 1.4, [body("BEFORE", 22, SEC, True, 2), code(ba["before"], 18)], fill=PANEL, line=LINE, anchor=MSO_ANCHOR.TOP)
    box(sl.shapes, 6.3, 3.5, 6.43, 1.4, [body("AFTER", 22, AMB, True, 2), code(ba["after"], 18)], fill=PANEL, line=AMB, lw=3, anchor=MSO_ANCHOR.TOP)
    textbox(sl.shapes, MX, 5.4, CW, 0.8, [body(sp["footer"], 32, TX, True)])


def k_appendix_steps(sp, sl):
    set_title(sl, sp["title"])
    for i, it in enumerate(sp["items"]):
        box(sl.shapes, MX, 2.0 + i * 1.3, 0.8, 0.8, [body(str(i + 1), 32, BG, True, align=PP_ALIGN.CENTER)], fill=TX, line=None, shape=MSO_SHAPE.OVAL, margin=0.02)
        textbox(sl.shapes, MX + 1.1, 2.0 + i * 1.3, 10.5, 0.8, [body(it, 36, TX, True)], anchor=MSO_ANCHOR.MIDDLE)


def k_appendix_cost(sp, sl):
    set_title(sl, sp["title"])
    for i, it in enumerate(sp["items"]):
        box(sl.shapes, MX, 2.0 + i * 1.3, CW, 1.05, [body(it, 32, TX, True)], fill=PANEL, line=LINE)


def k_appendix_qa(sp, sl):
    set_title(sl, sp["title"])
    for i, (q, a) in enumerate(sp["qa"]):
        y = 1.9 + i * 1.2
        textbox(sl.shapes, MX, y, 5.6, 1.0, [body(q, 28, TX, True)], anchor=MSO_ANCHOR.MIDDLE)
        textbox(sl.shapes, MX + 5.8, y, 6.3, 1.0, [body(a, 28, SEC)], anchor=MSO_ANCHOR.MIDDLE)


KINDS = {"title": k_title, "two_cards": k_two_cards, "panels": k_panels, "statement": k_statement, "hub": k_hub, "trace_reveal": k_trace_reveal,
         "json_compare": k_json_compare, "steps_trace": k_steps_trace, "handoff": k_handoff, "chart_retention": k_chart_retention, "architecture": k_architecture,
         "metrics_vs_drift": k_metrics_vs_drift, "demo_green": k_demo_green, "demo_catch": k_demo_catch, "demo_fix": k_demo_fix, "pattern_pin": k_pattern_pin,
         "pattern_memory": k_pattern_memory, "pattern_rules": k_pattern_rules, "pattern_replay": k_pattern_replay, "results_table": k_results_table,
         "checklist": k_checklist, "closing": k_closing, "appendix_text": k_appendix_text, "appendix_steps": k_appendix_steps, "appendix_cost": k_appendix_cost,
         "appendix_qa": k_appendix_qa}

FADE = {6, 7, 8, 9, 13, 14, 15, 16, 17, 18, 19}



NOTE_PNG = {10: "explorer_retention_95_3.png", 13: "explorer_green_day30.png", 14: "explorer_catch_day14.png", 15: "explorer_fix_all.png", 20: "explorer_results.png"}


def add_notes_picture(notes_slide, path):
    """Puts the backup PNG of the explorer scene on the notes page (python-pptx has no API for it)."""
    from pptx.oxml.shapes.picture import CT_Picture
    part = notes_slide.part
    image_part, rId = part.get_or_add_image_part(path)
    w, h = Inches(3.4), Inches(3.4 * 720 / 1280)
    pic = CT_Picture.new_pic(900, "Backup: explorer scene", "Backup PNG of the explorer scene for this slide", rId, Inches(3.8), Inches(8.2), w, h)
    notes_slide.shapes._spTree.append(pic)



def mmss(t):
    m, s = t.split(":")
    return int(m) * 60 + int(s)


def fmt(sec):
    return f"{sec // 60:02d}:{sec % 60:02d}"


mains = Y["slides"]
for i, (sp, sl) in enumerate(SLIDES):
    KINDS[sp["kind"]](sp, sl)
    if sp.get("results"):
        tag(sl)
    if i < len(mains) and int(sp["id"]) in FADE:
        fade_transition(sl)
    stamp_footer(sl, ("A" + str(sp["id"])[1:]) if str(sp["id"]).startswith("A") else None)
    # notes
    if i < len(mains):
        start = mmss(sp["start"])
        end = mmss(mains[i + 1]["start"]) if i + 1 < len(mains) else mmss(Y["meta"]["session_end"])
        stamp = f"[{fmt(start)} to {fmt(end)}]"
    else:
        stamp = f"[{fmt(mmss(Y['meta']['session_end']))} to {fmt(mmss(Y['meta']['qa_end']))}]"
    n = sp["notes"]
    words = len(n["script"].split())
    assert words <= 90, (sp["id"], words)
    cue = n["cue"]
    url = None
    for key in ("link", "links"):
        v = sp.get(key)
        if v:
            url = v["url"] if isinstance(v, dict) else v[0]["url"]
    short = n["short"]
    text = f"{stamp}\n{n['script']}\nCUE: {cue}\nIF SHORT ON TIME: {short}"
    sl.notes_slide.notes_text_frame.text = text
    png = NOTE_PNG.get(i + 1) if i < len(mains) else None
    if png and os.path.exists(os.path.join(ROOT, "export", png)):
        add_notes_picture(sl.notes_slide, os.path.join(ROOT, "export", png))

# the main slide timings must sum to 25:00
total = mmss(Y["meta"]["session_end"]) - mmss(mains[0]["start"])
assert total == 25 * 60

# template slides: welcome and bio before the talk, thank you after slide 22; the template's empty topic and questions slides are replaced by slides 1 and 22
def template_text(sl, old, new):
    for sh in sl.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)


def set_notes(sl, text):
    sl.notes_slide.notes_text_frame.text = text


welcome, bio, topic, questions, thanks = TPL
for sl in (bio, thanks):
    template_text(sl, "2025", "2026")
template_text(bio, "<Speaker Bio>", "[Speaker name, role]")
bs = [sh for sh in bio.shapes if sh.has_text_frame and "Speaker name" in sh.text_frame.text][0]
bs.top = Inches(0.55)
sty = bs.text_frame.paragraphs[0].runs[0].font
bb = bio.shapes.add_textbox(Inches(0.34), Inches(1.5), Inches(8.2), Inches(2.6))
fill_tf(bb.text_frame, [body("[One line on what you do]", 26, SEC, after=10), body("[One line on your AI and AWS work]", 26, SEC, after=10), body("[One line on community or teaching]", 26, SEC)])
template_text(thanks, "<Name>", "[Speaker name]")
template_text(thanks, "<linkedin handler>", "[linkedin handle]")
set_notes(welcome, "[before 0:00]\nEvent welcome slide from the template. Leave it up while people settle.\nCUE: none\nIF SHORT ON TIME: skip")
set_notes(bio, "[before 0:00]\nAbout 15 seconds. Who you are and why you care about this problem.\nCUE: none\nIF SHORT ON TIME: skip if the host has already introduced you")
set_notes(thanks, "[25:00 to 30:00]\nThank you. Questions are open. The appendix slides after this one cover other sources of drift, the six step checklist, cost and likely questions.\nCUE: none\nIF SHORT ON TIME: none")
# order: welcome, bio, slides 1 to 22, thank you, appendix
ids = list(prs.slides._sldIdLst)
by_slide = {sl.slide_id: el for el, sl in zip(ids, prs.slides)}
main = [by_slide[sl.slide_id] for sp, sl in SLIDES[:len(mains)]]
app = [by_slide[sl.slide_id] for sp, sl in SLIDES[len(mains):]]
order = [by_slide[welcome.slide_id], by_slide[bio.slide_id]] + main + [by_slide[thanks.slide_id]] + app
drop = [by_slide[t.slide_id].rId for t in (topic, questions)]
lst = prs.slides._sldIdLst
for el in list(lst):
    lst.remove(el)
for el in order:
    lst.append(el)
for rid in drop:
    prs.part.drop_rel(rid)

out = os.path.join(ROOT, "deck", "drift-talk.pptx")
import datetime
prs.core_properties.created = prs.core_properties.modified = datetime.datetime.utcnow()  # python-pptx keeps the template's dates otherwise
prs.core_properties.title = SHORT
prs.save(out)
print(f"deck/drift-talk.pptx written: {len(prs.slides)} slides ({len(mains)} main + {len(Y['appendix'])} appendix + 3 template), source={R['meta']['source']}")
